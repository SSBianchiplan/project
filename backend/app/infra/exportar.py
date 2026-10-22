"""Exportação para XLSX (openpyxl) e PDF (reportlab).

Um relatório é uma lista de seções; cada seção tem título, opcionalmente
uma lista de pares (rótulo, valor) e uma tabela (cabeçalhos + linhas) ou uma
série para gráfico. O mesmo modelo alimenta os dois formatos, então o que se
lê no PDF é o que se abre no Excel.
"""
from __future__ import annotations

import io
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.charts.lineplots import LinePlot
from reportlab.graphics.shapes import Drawing, Line, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

LOGO = Path(__file__).resolve().parents[1] / "estaticos" / "logo.png"
GRAFITE = colors.HexColor("#2d3336")
ACO = colors.HexColor("#7d8286")
CLARO = colors.HexColor("#edeff0")
LINHA = colors.HexColor("#dfe3e4")
EM_DIA = colors.HexColor("#2f7d4f")
ACAO = colors.HexColor("#a11d1d")
BASE = colors.HexColor("#b9bfc2")


@dataclass
class Serie:
    nome: str
    pontos: list[tuple[date, Decimal | None]]
    cor: str = "#7d8286"
    tracejada: bool = False


@dataclass
class Grafico:
    """linha: séries temporais (curva S, custo). barras: grupos (rótulo → [(série, valor, cor)]).
    gantt: barras horizontais por atividade (rótulo, início, término, crítica, cor, avanço 0-1)."""
    titulo: str
    series: list[Serie] = field(default_factory=list)
    formato_y: str = "pct"          # pct | moeda | numero
    tipo: str = "linha"             # linha | barras | gantt
    barras: list[tuple[str, list[tuple[str, Decimal, str]]]] = field(default_factory=list)
    gantt: list[tuple[str, date, date, bool, str, Decimal]] = field(default_factory=list)
    marcador: date | None = None    # linha vertical (data de status) no gantt
    altura_mm: float = 62


@dataclass
class Secao:
    titulo: str
    cartoes: list[tuple[str, str]] = field(default_factory=list)      # (rótulo, valor)
    cabecalhos: list[str] = field(default_factory=list)
    linhas: list[list[Any]] = field(default_factory=list)
    grafico: Grafico | None = None
    nota: str = ""
    larguras: list[float] | None = None     # proporções das colunas no PDF


@dataclass
class Cabecalho:
    titulo: str
    contrato_codigo: str
    empreendimento: str
    contratante: str
    contratada: str
    data_corte: date
    emitido_por: str
    baseline: str = ""
    paisagem: bool = True


@dataclass
class Relatorio:
    cabecalho: Cabecalho
    secoes: list[Secao]


# ------------------------------------------------------------------ formatos

def fmt(v: Any) -> str:
    if v is None:
        return "—"
    if isinstance(v, bool):
        return "sim" if v else "não"
    if isinstance(v, Decimal):
        return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if isinstance(v, float):
        return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return str(v)


def pct(v: Decimal | float | None, casas: int = 1) -> str:
    return "—" if v is None else f"{float(v) * 100:.{casas}f}%".replace(".", ",")


def moeda(v: Decimal | float | None) -> str:
    return "—" if v is None else ("R$ " + fmt(Decimal(str(v)).quantize(Decimal("0.01"))))


# --------------------------------------------------------------------- XLSX

def xlsx(rel: Relatorio) -> bytes:
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Capa"
    c = rel.cabecalho
    negrito = Font(bold=True)
    for i, (k, v) in enumerate([
        ("Relatório", c.titulo), ("Sistema", "GPlanSiS|Project"),
        ("Contrato", c.contrato_codigo), ("Empreendimento", c.empreendimento), ("Contratante", c.contratante),
        ("Contratada", c.contratada), ("Linha de base", c.baseline or "—"), ("Data de corte", fmt(c.data_corte)),
        ("Emitido por", c.emitido_por), ("Emitido em", fmt(datetime.now())),
    ], start=1):
        ws0.cell(row=i, column=1, value=k).font = negrito
        ws0.cell(row=i, column=2, value=v)
    ws0.column_dimensions["A"].width = 18
    ws0.column_dimensions["B"].width = 70

    cab_fill = PatternFill("solid", fgColor="2D3336")
    cab_font = Font(bold=True, color="F7F8F8")
    fino = Side(style="thin", color="DFE3E4")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    usados: set[str] = {"Capa"}
    for sec in rel.secoes:
        nome = sec.titulo[:28].replace("/", "-").replace(":", "")
        base_nome, k = nome, 2
        while nome in usados:
            nome, k = f"{base_nome[:25]} {k}", k + 1
        usados.add(nome)
        ws = wb.create_sheet(nome)
        linha = 1
        ws.cell(row=linha, column=1, value=sec.titulo).font = Font(bold=True, size=13)
        linha += 2
        for k_, v in sec.cartoes:
            ws.cell(row=linha, column=1, value=k_).font = negrito
            ws.cell(row=linha, column=2, value=v)
            linha += 1
        if sec.cartoes:
            linha += 1
        if sec.grafico:
            ws.cell(row=linha, column=1, value=sec.grafico.titulo).font = negrito
            linha += 1
            ws.cell(row=linha, column=1, value="Semana").font = cab_font
            ws.cell(row=linha, column=1).fill = cab_fill
            for j, s_ in enumerate(sec.grafico.series, start=2):
                cel = ws.cell(row=linha, column=j, value=s_.nome)
                cel.font, cel.fill = cab_font, cab_fill
            linha += 1
            datas = sorted({d for s_ in sec.grafico.series for d, _ in s_.pontos})
            mapa = [{d: v for d, v in s_.pontos} for s_ in sec.grafico.series]
            for d in datas:
                ws.cell(row=linha, column=1, value=d)
                for j, m in enumerate(mapa, start=2):
                    v = m.get(d)
                    ws.cell(row=linha, column=j, value=None if v is None else float(v))
                linha += 1
            linha += 1
        if sec.cabecalhos:
            for j, h in enumerate(sec.cabecalhos, start=1):
                cel = ws.cell(row=linha, column=j, value=h)
                cel.font, cel.fill, cel.border = cab_font, cab_fill, borda
                cel.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = ws.cell(row=linha + 1, column=1)
            linha += 1
            for r in sec.linhas:
                for j, bruto in enumerate(r, start=1):
                    # Excel não aceita fuso: converte para hora local sem tzinfo
                    v = bruto.astimezone().replace(tzinfo=None) if isinstance(bruto, datetime) and bruto.tzinfo else bruto
                    cel = ws.cell(row=linha, column=j, value=float(v) if isinstance(v, Decimal) else v)
                    cel.border = borda
                    if isinstance(v, Decimal | float | int) and not isinstance(v, bool):
                        cel.number_format = "#,##0.00"
                        cel.alignment = Alignment(horizontal="right")
                    elif isinstance(v, date):
                        cel.number_format = "DD/MM/YYYY"
                linha += 1
            for j in range(1, len(sec.cabecalhos) + 1):
                larg = max([len(str(sec.cabecalhos[j - 1]))] + [len(fmt(r[j - 1])) for r in sec.linhas if j - 1 < len(r)][:200])
                ws.column_dimensions[get_column_letter(j)].width = min(max(10, larg + 2), 60)
        if sec.nota:
            linha += 1
            ws.cell(row=linha, column=1, value=sec.nota).font = Font(italic=True, color="7D8286")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------- PDF

def _estilos() -> dict[str, ParagraphStyle]:
    ss = getSampleStyleSheet()
    return {
        "titulo": ParagraphStyle("t", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=15, textColor=GRAFITE, spaceAfter=2),
        "sub": ParagraphStyle("s", parent=ss["Normal"], fontName="Helvetica", fontSize=8.5, textColor=ACO, leading=11),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, textColor=ACO,
                             spaceBefore=10, spaceAfter=4, textTransform="uppercase"),
        "celula": ParagraphStyle("c", parent=ss["Normal"], fontName="Helvetica", fontSize=7.6, leading=9.2, textColor=GRAFITE),
        "celula_dir": ParagraphStyle("cd", parent=ss["Normal"], fontName="Helvetica", fontSize=7.6, leading=9.2, alignment=TA_RIGHT, textColor=GRAFITE),
        "cab": ParagraphStyle("cb", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=7.6, leading=9.2, textColor=colors.white),
        "cartao_rot": ParagraphStyle("cr", parent=ss["Normal"], fontName="Helvetica", fontSize=7, textColor=ACO, leading=8.5),
        "cartao_val": ParagraphStyle("cv", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=11.5, textColor=GRAFITE, leading=14),
        "nota": ParagraphStyle("n", parent=ss["Normal"], fontName="Helvetica-Oblique", fontSize=7.5, textColor=ACO, leading=9.5, alignment=TA_LEFT),
    }


def _cortar(canvas, txt: str, fonte: str, tam: float, largura: float) -> str:  # noqa: ANN001
    if canvas.stringWidth(txt, fonte, tam) <= largura:
        return txt
    while txt and canvas.stringWidth(txt + "…", fonte, tam) > largura:
        txt = txt[:-1]
    return txt + "…"


def _cabecalho_rodape(cab: Cabecalho):  # noqa: ANN202 — assinatura exigida pelo reportlab
    def desenhar(canvas, doc) -> None:  # noqa: ANN001
        canvas.saveState()
        w, h = doc.pagesize
        m = 14 * mm
        y = h - 12 * mm
        if LOGO.exists():
            canvas.drawImage(str(LOGO), m, y - 9 * mm, width=11 * mm, height=11 * mm, mask="auto")
        # linha 1: sistema (esq.) e título do relatório (dir.)
        canvas.setFillColor(GRAFITE)
        canvas.setFont("Helvetica-Bold", 10.5)
        canvas.drawString(28 * mm, y - 2.2 * mm, "GPlanSiS|Project")
        canvas.setFont("Helvetica-Bold", 9)
        canvas.drawRightString(w - m, y - 2.2 * mm, cab.titulo)
        # linha 2: procedimento (esq.), contrato (dir.) — cada lado com metade da largura
        meio = (w - 2 * m) / 2
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(ACO)
        canvas.drawString(28 * mm, y - 6 * mm, _cortar(canvas, "Planejamento e controle de projetos · GPlanSiS|Group", "Helvetica", 7.5, meio - 14 * mm))
        canvas.drawRightString(w - m, y - 6 * mm, _cortar(canvas, f"{cab.contrato_codigo} · {cab.empreendimento}", "Helvetica", 7.5, meio))
        # linha 3: partes e corte, largura inteira, alinhada à direita
        detalhe = (f"Contratante {cab.contratante} · Contratada {cab.contratada}"
                   f"{(' · BL' + cab.baseline) if cab.baseline else ''} · data de corte {fmt(cab.data_corte)}")
        canvas.drawRightString(w - m, y - 9.5 * mm, _cortar(canvas, detalhe, "Helvetica", 7.5, w - 2 * m - 16 * mm))
        canvas.setStrokeColor(LINHA)
        canvas.setLineWidth(0.6)
        canvas.line(m, y - 11.5 * mm, w - m, y - 11.5 * mm)
        # rodapé
        canvas.line(m, 12 * mm, w - m, 12 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.drawString(m, 8 * mm, _cortar(canvas, f"Emitido por {cab.emitido_por} em {fmt(datetime.now())} · documento gerado pelo sistema; números reproduzíveis pela memória de cálculo", "Helvetica", 7, w - 2 * m - 30 * mm))
        canvas.drawRightString(w - m, 8 * mm, f"página {doc.page}")
        canvas.restoreState()
    return desenhar


def _grafico(g: Grafico, largura: float) -> Drawing:
    if g.tipo == "barras":
        return _grafico_barras(g, largura)
    if g.tipo == "gantt":
        return _grafico_gantt(g, largura)
    alt = 62 * mm
    d = Drawing(largura, alt)
    lp = LinePlot()
    lp.x, lp.y, lp.width, lp.height = 14 * mm, 8 * mm, largura - 20 * mm, alt - 16 * mm
    datas = sorted({dt for s in g.series for dt, _ in s.pontos})
    if not datas:
        return d
    x0 = datas[0].toordinal()
    dados = []
    for s in g.series:
        pts = [((dt.toordinal() - x0) / 7.0, float(v)) for dt, v in s.pontos if v is not None]
        dados.append(pts or [(0, 0)])
    lp.data = dados
    for i, s in enumerate(g.series):
        lp.lines[i].strokeColor = colors.HexColor(s.cor)
        lp.lines[i].strokeWidth = 1.6 if i else 1.2
        if s.tracejada:
            lp.lines[i].strokeDashArray = [3, 2]
    maxy = max([v for pts in dados for _, v in pts] + [0.0001])
    lp.yValueAxis.valueMin = 0
    lp.yValueAxis.valueMax = 1.0 if g.formato_y == "pct" and maxy <= 1.0001 else maxy * 1.05
    if g.formato_y == "pct":
        lp.yValueAxis.labelTextFormat = lambda v: f"{v * 100:.0f}%"
    elif g.formato_y == "numero":
        lp.yValueAxis.labelTextFormat = lambda v: f"{v:,.0f}".replace(",", ".")
    else:
        lp.yValueAxis.labelTextFormat = lambda v: f"{v / 1000:,.0f} mil".replace(",", ".")
    lp.yValueAxis.labels.fontName = "Helvetica"
    lp.yValueAxis.labels.fontSize = 6.5
    lp.yValueAxis.strokeColor = LINHA
    lp.yValueAxis.gridStrokeColor = LINHA
    lp.yValueAxis.visibleGrid = 1
    lp.xValueAxis.valueMin = 0
    lp.xValueAxis.valueMax = (datas[-1].toordinal() - x0) / 7.0
    lp.xValueAxis.labels.fontName = "Helvetica"
    lp.xValueAxis.labels.fontSize = 6.5
    lp.xValueAxis.strokeColor = LINHA
    passo = max(1, len(datas) // 12)
    lp.xValueAxis.valueSteps = [(dt.toordinal() - x0) / 7.0 for dt in datas[::passo]]
    lp.xValueAxis.labelTextFormat = lambda v: (date.fromordinal(int(round(v * 7)) + x0)).strftime("%d/%m")
    d.add(lp)
    leg = Legend()
    leg.x, leg.y = 16 * mm, alt - 4 * mm
    leg.fontName, leg.fontSize = "Helvetica", 6.8
    leg.alignment = "right"
    leg.columnMaximum = 1
    leg.colorNamePairs = [(colors.HexColor(s.cor), s.nome) for s in g.series]
    d.add(leg)
    d.add(String(largura - 2 * mm, alt - 4 * mm, g.titulo, fontName="Helvetica-Bold", fontSize=7.5, fillColor=ACO, textAnchor="end"))
    return d


def _fmt_y(g: Grafico) -> Callable[[float], str]:
    if g.formato_y == "pct":
        return lambda v: f"{v * 100:.0f}%"
    if g.formato_y == "numero":
        return lambda v: f"{v:,.0f}".replace(",", ".")
    return lambda v: f"{v / 1000:,.0f} mil".replace(",", ".")


def _grafico_barras(g: Grafico, largura: float) -> Drawing:
    """Barras agrupadas: cada grupo é um rótulo no eixo X; cada série tem cor própria."""
    alt = g.altura_mm * mm
    d = Drawing(largura, alt)
    if not g.barras:
        return d
    nomes = []
    for _, itens in g.barras:
        for nome, _, _ in itens:
            if nome not in nomes:
                nomes.append(nome)
    cores = {}
    for _, itens in g.barras:
        for nome, _, cor in itens:
            cores.setdefault(nome, cor)
    bc = VerticalBarChart()
    bc.x, bc.y, bc.width, bc.height = 14 * mm, 9 * mm, largura - 20 * mm, alt - 18 * mm
    bc.data = [[float(next((v for n, v, _ in itens if n == nome), 0)) for _, itens in g.barras] for nome in nomes]
    bc.categoryAxis.categoryNames = [r for r, _ in g.barras]
    bc.categoryAxis.labels.fontName, bc.categoryAxis.labels.fontSize = "Helvetica", 6.5
    bc.categoryAxis.labels.angle = 30 if len(g.barras) > 8 else 0
    bc.categoryAxis.labels.dy = -8 if len(g.barras) > 8 else -2
    bc.categoryAxis.strokeColor = LINHA
    maxy = max([v for fila in bc.data for v in fila] + [0.0001])
    bc.valueAxis.valueMin = 0
    bc.valueAxis.valueMax = 1.0 if g.formato_y == "pct" and maxy <= 1.0001 else maxy * 1.12
    bc.valueAxis.labelTextFormat = _fmt_y(g)
    bc.valueAxis.labels.fontName, bc.valueAxis.labels.fontSize = "Helvetica", 6.5
    bc.valueAxis.strokeColor = LINHA
    bc.valueAxis.gridStrokeColor = LINHA
    bc.valueAxis.visibleGrid = 1
    bc.groupSpacing = 6
    bc.barSpacing = 1
    for i, nome in enumerate(nomes):
        bc.bars[i].fillColor = colors.HexColor(cores[nome])
        bc.bars[i].strokeColor = colors.white
        bc.bars[i].strokeWidth = 0.3
    d.add(bc)
    if len(nomes) > 1:
        leg = Legend()
        leg.x, leg.y = 16 * mm, alt - 4 * mm
        leg.fontName, leg.fontSize = "Helvetica", 6.8
        leg.alignment = "right"
        leg.columnMaximum = 1
        leg.colorNamePairs = [(colors.HexColor(cores[n]), n) for n in nomes]
        d.add(leg)
    d.add(String(largura - 2 * mm, alt - 4 * mm, g.titulo, fontName="Helvetica-Bold", fontSize=7.5, fillColor=ACO, textAnchor="end"))
    return d


def _grafico_gantt(g: Grafico, largura: float) -> Drawing:
    """Gantt compacto: uma linha por atividade (rótulo à esquerda), barra colorida por fase, crítica em vermelho,
    avanço como sombra interna, marcador tracejado na data de status."""
    linhas = g.gantt
    if not linhas:
        return Drawing(largura, 10 * mm)
    lh = 3.6 * mm
    top = 7 * mm
    alt = top + lh * len(linhas) + 8 * mm
    d = Drawing(largura, alt)
    x_rot = 46 * mm
    x0, x1 = x_rot + 2 * mm, largura - 3 * mm
    ini = min(i for _, i, _, _, _, _ in linhas)
    fim = max(f for _, _, f, _, _, _ in linhas)
    if g.marcador:
        ini, fim = min(ini, g.marcador), max(fim, g.marcador)
    span = max((fim - ini).days, 1)
    px = (x1 - x0) / span
    # eixo de meses
    m = date(ini.year, ini.month, 1)
    while m <= fim:
        x = x0 + max((m - ini).days, 0) * px
        d.add(Line(x, alt - top + 1 * mm, x, 6 * mm, strokeColor=LINHA, strokeWidth=0.4))
        d.add(String(x + 1, alt - top + 2.2 * mm, m.strftime("%m/%y"), fontName="Helvetica", fontSize=5.8, fillColor=ACO))
        m = date(m.year + (m.month // 12), m.month % 12 + 1, 1)
    for k, (rot, i, f, crit, cor, av) in enumerate(linhas):
        y = alt - top - (k + 1) * lh + 0.6 * mm
        if k % 2 == 0:
            d.add(Rect(2 * mm, y - 0.5 * mm, largura - 4 * mm, lh, fillColor=colors.HexColor("#f7f8f8"), strokeColor=None))
        d.add(String(2.5 * mm, y + 0.5 * mm, rot.lstrip("*")[:36], fontName="Helvetica-Bold" if rot.startswith("*") else "Helvetica", fontSize=5.8, fillColor=GRAFITE))
        xa = x0 + (i - ini).days * px
        w = max(((f - i).days + 1) * px, 1.2)
        if i == f and rot.strip().endswith("◆"):
            d.add(Rect(xa - 1.2, y + 0.2 * mm, 2.4, 2.4, fillColor=GRAFITE, strokeColor=None))
            continue
        d.add(Rect(xa, y, w, lh - 1.2 * mm, fillColor=colors.HexColor("#a11d1d") if crit else colors.HexColor(cor), strokeColor=None))
        if av > 0:
            d.add(Rect(xa, y, w * float(min(av, 1)), lh - 1.2 * mm, fillColor=colors.Color(0, 0, 0, alpha=0.32), strokeColor=None))
    if g.marcador:
        xm = x0 + (g.marcador - ini).days * px
        d.add(Line(xm, 5 * mm, xm, alt - top + 1 * mm, strokeColor=GRAFITE, strokeWidth=0.8, strokeDashArray=[2, 1.5]))
        d.add(String(xm + 1, 2.5 * mm, f"status {g.marcador.strftime('%d/%m/%Y')}", fontName="Helvetica", fontSize=5.8, fillColor=GRAFITE))
    d.add(String(largura - 2 * mm, 2.5 * mm, g.titulo, fontName="Helvetica-Bold", fontSize=6.5, fillColor=ACO, textAnchor="end"))
    return d


def pdf(rel: Relatorio) -> bytes:
    cab = rel.cabecalho
    pagesize = landscape(A4) if cab.paisagem else A4
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=pagesize, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=26 * mm, bottomMargin=16 * mm, title=cab.titulo, author=cab.emitido_por,
                            subject=f"{cab.contrato_codigo} — {cab.empreendimento}")
    st = _estilos()
    largura = pagesize[0] - 28 * mm
    fl: list[Any] = [Paragraph(cab.titulo, st["titulo"]),
                     Paragraph(f"{cab.contrato_codigo} · {cab.empreendimento} · data de corte {fmt(cab.data_corte)}", st["sub"]),
                     Spacer(1, 4)]
    for sec in rel.secoes:
        bloco: list[Any] = [Paragraph(sec.titulo, st["h2"])]
        if sec.cartoes:
            n = len(sec.cartoes)
            por_linha = 6 if cab.paisagem else 4
            for i in range(0, n, por_linha):
                fatia = sec.cartoes[i:i + por_linha]
                celulas = [[Paragraph(k, st["cartao_rot"]) for k, _ in fatia], [Paragraph(v, st["cartao_val"]) for _, v in fatia]]
                t = Table(celulas, colWidths=[largura / por_linha] * len(fatia))
                t.setStyle(TableStyle([
                    ("BOX", (0, 0), (-1, -1), 0.5, LINHA), ("INNERGRID", (0, 0), (-1, -1), 0.5, LINHA),
                    ("BACKGROUND", (0, 0), (-1, -1), colors.white), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, 0), 4), ("BOTTOMPADDING", (0, 1), (-1, 1), 5),
                    ("LINEBELOW", (0, 0), (-1, 0), 0, colors.white)]))
                bloco.append(t)
                bloco.append(Spacer(1, 3))
        if sec.grafico and sec.grafico.tipo == "gantt":
            # gantt longo: fatias de até 24 linhas, cada uma cabe numa página paisagem
            fl.extend(bloco)
            linhas_g = sec.grafico.gantt
            for i in range(0, len(linhas_g), 24):
                fatia = Grafico(sec.grafico.titulo if i == 0 else f"{sec.grafico.titulo} (continuação)", tipo="gantt", gantt=linhas_g[i:i + 24], marcador=sec.grafico.marcador)
                fl.append(_grafico(fatia, largura - 14))
                fl.append(Spacer(1, 4))
            if sec.nota:
                fl.append(Paragraph(sec.nota, st["nota"]))
            fl.append(Spacer(1, 6))
            continue
        if sec.grafico:
            bloco.append(_grafico(sec.grafico, largura - 14))
            bloco.append(Spacer(1, 4))
        if sec.cabecalhos:
            dados = [[Paragraph(h, st["cab"]) for h in sec.cabecalhos]]
            numericas = [all(isinstance(r[j], Decimal | float | int) and not isinstance(r[j], bool) for r in sec.linhas if j < len(r) and r[j] is not None)
                         for j in range(len(sec.cabecalhos))]
            for r in sec.linhas:
                dados.append([Paragraph(fmt(v), st["celula_dir"] if numericas[j] else st["celula"]) for j, v in enumerate(r)])
            props = sec.larguras or [1.0] * len(sec.cabecalhos)
            soma = sum(props)
            t = Table(dados, colWidths=[largura * p / soma for p in props], repeatRows=1)
            estilo = [("BACKGROUND", (0, 0), (-1, 0), GRAFITE), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                      ("LINEBELOW", (0, 0), (-1, -1), 0.4, LINHA), ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                      ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
            for i in range(1, len(dados)):
                if i % 2 == 0:
                    estilo.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#f7f8f8")))
            t.setStyle(TableStyle(estilo))
            bloco.append(t)
        if sec.nota:
            bloco.append(Spacer(1, 3))
            bloco.append(Paragraph(sec.nota, st["nota"]))
        # tabela longa flui entre páginas (cabeçalho repete); bloco curto fica junto
        if sec.cabecalhos and len(sec.linhas) > 8:
            fl.extend(bloco)
        else:
            fl.append(KeepTogether(bloco))
        fl.append(Spacer(1, 6))
    if fl and isinstance(fl[-1], PageBreak):
        fl.pop()
    desenhar = _cabecalho_rodape(cab)
    doc.build(fl, onFirstPage=desenhar, onLaterPages=desenhar)
    return buf.getvalue()
