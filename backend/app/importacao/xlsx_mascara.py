"""Máscara GPS-PLN-ANX-IV (planilha). Formato de entrada mais confiável, porque
as colunas são as do procedimento e não dependem de tradução de campo."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from .modelo import AtividadeImportada, Formato, NoImportada, PlanoImportado, VinculoImportado

ABA = "Cronograma"
LINHA_CABECALHO = 4


def _data(v: object) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def parse(conteudo: bytes, origem: str = "") -> PlanoImportado:
    wb = load_workbook(BytesIO(conteudo), data_only=True, read_only=True)
    if ABA not in wb.sheetnames:
        raise ValueError(f"planilha sem a aba '{ABA}' — não é a máscara GPS-PLN-ANX-IV")
    ws = wb[ABA]
    plano = PlanoImportado(formato=Formato.XLSX_MASCARA, origem=origem)

    linhas = list(ws.iter_rows(min_row=LINHA_CABECALHO + 1, values_only=True))
    for linha in linhas:
        if not linha or linha[0] is None:
            continue
        ident, edt, nivel, nome, tipo, dur, inicio, termino, pred, restr = linha[:10]
        edt_s = str(edt).strip()
        if tipo == "Resumo":
            plano.nos.append(NoImportada(edt_s, str(nome).strip(), int(nivel)))
            continue
        marco = tipo == "Marco"
        plano.atividades.append(AtividadeImportada(
            id_externo=str(ident),
            edt=edt_s,
            nome=str(nome).strip(),
            duracao_dias=0 if marco else int(Decimal(str(dur or 0))),
            inicio=_data(inicio),
            termino=_data(termino),
            marco=marco,
            restricao=str(restr).strip() if restr else None,
        ))
        if pred:
            for bruto in str(pred).split(";"):
                vinculo = _parse_predecessora(bruto, str(ident))
                if vinculo is not None:
                    plano.vinculos.append(vinculo)
    return plano


def _parse_predecessora(bruto: str, sucessora: str) -> VinculoImportado | None:
    """'12TI+3d' -> vínculo TI com lag 3. Lag negativo é preservado: quem
    recusa é o validador, com a invariante nomeada, não o parser em silêncio."""
    item = bruto.strip()
    if not item:
        return None
    tipo = "TI"
    for marca in ("TI", "II", "TT", "IT"):
        if marca in item:
            tipo = marca
            item = item.replace(marca, "")
            break
    lag = 0
    if "+" in item:
        item, lag_txt = item.split("+", 1)
        lag = int(Decimal(lag_txt.strip().replace("d", "") or 0))
    elif "-" in item[1:]:
        corpo, lag_txt = item[1:].split("-", 1)
        item = item[0] + corpo
        lag = -int(Decimal(lag_txt.strip().replace("d", "") or 0))
    return VinculoImportado(item.strip(), sucessora, tipo, lag)
